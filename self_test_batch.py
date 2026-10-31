#!/usr/bin/env python3
# =============================================================
# АВТОНОМНАЯ САМОПРОВЕРКА ПАКЕТНОЙ ПЕЧАТИ
# =============================================================
"""Проверяет пакетную печать без реальных заявок и без Telegram.

Сценарий:
  1. создаёт временную SQLite-базу и 10 тестовых заявок;
  2. формирует два Word-файла (перед/зад) через боевой код;
  3. запускает подтверждение пакета с заглушкой Telegram;
  4. проверяет строки реестра 25–16, номера ТТН и открытость пачки;
  5. проверяет отсутствие пустых страниц через LibreOffice/PDF;
  6. оставляет образцы Word/PDF и отчёт в «Результат/Самопроверка...».

Реальные orders.db, data/ и Результат/боевых заявок не изменяются.
"""
import argparse
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

PROJECT_ROOT = Path(__file__).resolve().parent
MASTER_ID = 999001
SELFTEST_HOME = Path(tempfile.mkdtemp(prefix="ttn_batch_selftest_home_"))
os.environ.setdefault("BOT_TOKEN", "123456:SELF_TEST_TOKEN")
os.environ["YOUR_TELEGRAM_ID"] = str(MASTER_ID)
# config.py при старте сохраняет masters.json и last_orders.json.
# Направляем их во временную папку, чтобы самопроверка не меняла production.
os.environ["MASTERS_FILE"] = str(SELFTEST_HOME / "masters.json")
os.environ["LAST_ORDERS_FILE"] = str(SELFTEST_HOME / "last_orders.json")
os.environ["TTN_DB_FILE"] = str(SELFTEST_HOME / "orders.db")
os.environ.setdefault("TTN_LOG_TO_FILE", "0")
sys.path.insert(0, str(PROJECT_ROOT))
os.chdir(PROJECT_ROOT)

from docx import Document

import batch_print
import batch_word
import config
import orders_db
import registry_filler
import utils
import word_filler

START_NUMBER = "7777000"
ORDER_COUNT = 10


def _order_data(number):
    quantity = number
    price = 1000
    total_cost = price * quantity
    vat_cost = round(total_cost * 0.2, 2)
    total_with_vat = round(total_cost + vat_cost, 2)
    vat_text, vat_cents = utils.amount_to_words_parts(vat_cost)
    total_text, total_cents = utils.amount_to_words_parts(total_with_vat)

    return {
        "day": "22",
        "month": "10",
        "year": "26",
        "waybill_number": f"САМОТЕСТ-{number:02d}",
        "car_model": "МАЗ-САМОТЕСТ",
        "car_plate": f"TEST {number:02d}",
        "trailer_info": "-",
        "loading_point": "Склад самопроверки",
        "unloading_point": f"Объект {number}",
        "driver_name": f"Тестовый водитель {number}",
        "total_vat_text": vat_text,
        "total_vat_cents": vat_cents,
        "total_with_vat_text": total_text,
        "total_with_vat_cents": total_cents,
        "cargos": [{
            "cargo_name": f"Тестовый груз {number}",
            "si": "т",
            "unit": "т",
            "quantity": str(quantity),
            "cargo_qty": str(quantity),
            "places": str(quantity),
            "price": str(price),
            "total_cost": f"{total_cost:.2f}",
            "vat_rate": "20%",
            "vat_cost": f"{vat_cost:.2f}",
            "total_with_vat": f"{total_with_vat:.2f}",
        }],
    }


class Report:
    def __init__(self):
        self.lines = []
        self.failures = 0
        self.skips = 0

    def check(self, condition, label, detail=""):
        if condition:
            self.lines.append(f"[OK]   {label}")
        else:
            self.failures += 1
            self.lines.append(f"[FAIL] {label}" + (f" — {detail}" if detail else ""))
        return bool(condition)

    def skip(self, label):
        self.skips += 1
        self.lines.append(f"[SKIP] {label}")

    def info(self, text):
        self.lines.append(text)

    def text(self):
        return "\n".join(self.lines)


def find_template(explicit):
    candidates = []
    if explicit:
        candidates.append(Path(explicit))
    candidates.extend([
        PROJECT_ROOT / "Шаблоны" / "Шаблон_ТТН.docx",
        PROJECT_ROOT / "assets" / "Шаблон_ТТН_демо.docx",
    ])
    for path in candidates:
        if path and path.exists():
            return path
    raise FileNotFoundError(
        "Не найден шаблон ТТН. Укажите его через --template.")


def find_soffice():
    found = shutil.which("soffice") or shutil.which("soffice.exe")
    if found:
        return found
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
    ):
        if os.path.exists(candidate):
            return candidate
    return None


def _section_layout(doc):
    return [
        (section.page_width, section.page_height,
         section.left_margin, section.right_margin,
         section.top_margin, section.bottom_margin,
         section.orientation)
        for section in doc.sections
    ]


def _document_text(doc):
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _copy_outputs(record, registry_path, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    front = output_dir / "Пакет_10_заявок_передние.docx"
    back = output_dir / "Пакет_10_заявок_задние.docx"
    registry = output_dir / "Реестр_10_заявок.docx"
    shutil.copy2(record["front_path"], front)
    shutil.copy2(record["back_path"], back)
    shutil.copy2(registry_path, registry)
    return front, back, registry


def _convert_to_pdf(soffice, path, out_dir):
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf",
         "--outdir", str(out_dir), str(path)],
        check=True, capture_output=True, timeout=180,
    )
    return Path(path).with_suffix(".pdf")


def run_selftest(template, output_dir, keep_temp=False):
    report = Report()
    runtime = Path(tempfile.mkdtemp(prefix="ttn_batch_selftest_"))
    report.info("=" * 70)
    report.info("САМОПРОВЕРКА ПАКЕТНОЙ ПЕЧАТИ ТТН")
    report.info(f"Дата: {datetime.now():%d.%m.%Y %H:%M:%S}")
    report.info(f"Шаблон: {template}")
    report.info(f"Временная папка: {runtime}")
    report.info(f"Временная конфигурация: {SELFTEST_HOME}")
    report.info(f"Папка результатов: {output_dir}")
    report.info("=" * 70)

    try:
        # Изолированная база и папки: боевые данные не затрагиваются.
        report.check(
            Path(config.MASTERS_FILE).resolve().is_relative_to(
                SELFTEST_HOME.resolve())
            and Path(config.LAST_ORDERS_FILE).resolve().is_relative_to(
                SELFTEST_HOME.resolve()),
            "masters.json и last_orders.json не затрагивают production")
        report.check(
            Path(orders_db.DB_FILE).resolve().is_relative_to(
                SELFTEST_HOME.resolve()),
            "orders.db не затрагивает production")
        orders_db.DB_FILE = str(runtime / "orders.db")
        orders_db.init_db()
        batch_print.PENDING_FILE = str(runtime / "pending.json")
        batch_print._selection_state.clear()
        word_filler.TEMPLATE = str(template)
        word_filler.RESULT_FOLDER = str(runtime / "Результат")
        registry_filler.RESULT_FOLDER = str(runtime / "Реестр")

        order_ids = []
        for number in range(1, ORDER_COUNT + 1):
            order_ids.append(orders_db.save_order(
                _order_data(number),
                master_id=MASTER_ID,
                driver_id=900000 + number,
            ))
        report.check(
            len(order_ids) == ORDER_COUNT,
            f"Создано тестовых заявок: {len(order_ids)}")

        # Контрольный документ нужен, чтобы сравнить размер бумаги и поля.
        source_path, _, _ = word_filler.fill_documents(
            _order_data(1), "selftest_source",
            output_dir=str(runtime / "source"))
        source_layout = _section_layout(Document(source_path))

        assignments = batch_print._make_assignments(
            order_ids, START_NUMBER, 0)
        plan = {
            "selected": order_ids,
            "start_number": START_NUMBER,
            "batch_id": None,
            "counter": 0,
            "assignments": assignments,
        }
        record = batch_print._generate_batch(MASTER_ID, plan)

        report.check(
            Path(record["front_path"]).exists()
            and Path(record["back_path"]).exists(),
            "Сформированы два Word-файла: передние и задние стороны")
        report.check(
            batch_print._get_pending(MASTER_ID) is not None,
            "Готовый пакет сохранён для подтверждения")

        front_doc = Document(record["front_path"])
        back_doc = Document(record["back_path"])
        expected_sections = ORDER_COUNT * batch_word.COPIES_PER_TTN
        report.check(
            len(front_doc.sections) == expected_sections,
            f"Передний файл: {len(front_doc.sections)} разделов "
            f"(ожидалось {expected_sections})")
        report.check(
            len(back_doc.sections) == expected_sections,
            f"Задний файл: {len(back_doc.sections)} разделов "
            f"(ожидалось {expected_sections})")
        report.check(
            _section_layout(front_doc) == [source_layout[0]] * expected_sections,
            "Передние страницы сохранили размер бумаги и поля шаблона")
        report.check(
            _section_layout(back_doc) == [source_layout[1]] * expected_sections,
            "Задние страницы сохранили размер бумаги и поля шаблона")
        for label, doc in (("переднем", front_doc), ("заднем", back_doc)):
            text = _document_text(doc)
            report.check(
                "{{" not in text and "cargo." not in text,
                f"В {label} файле нет незаполненных подстановок стоимости")

        # Подтверждение запускаем через настоящий обработчик, но Telegram
        # подменён заглушкой: наружу ничего не отправляется.
        messages = []

        def fake_send_message(*args, **kwargs):
            messages.append((args, kwargs))
            return SimpleNamespace(message_id=1)

        batch_print.bot.send_message = fake_send_message
        batch_print.bot.answer_callback_query = lambda *args, **kwargs: None
        call = SimpleNamespace(
            id="selftest-callback",
            data="batch_confirm_print",
            message=SimpleNamespace(
                chat=SimpleNamespace(id=MASTER_ID),
                message_id=1,
            ),
        )
        batch_print.batch_confirm_print(call)

        active = orders_db.get_active_batch(MASTER_ID)
        report.check(active is not None, "После подтверждения открыта пачка")
        batch_id = int(active["id"]) if active else None
        report.check(
            batch_id is not None and int(active["counter"]) == ORDER_COUNT,
            "Счётчик пачки увеличен на 10 заявок")
        report.check(
            batch_id is not None and active.get("closed_at") is None,
            "Пачка осталась открытой для добора заявок")
        report.check(
            batch_print._get_pending(MASTER_ID) is None,
            "После подтверждения ожидающий пакет снят")

        assigned = [
            orders_db.get_order(order_id) for order_id in order_ids
        ]
        report.check(
            all(order["status"] == "printed" for order in assigned),
            "Все 10 заявок переведены в статус «распечатана»")
        report.check(
            all(order["ttn_number"] for order in assigned),
            "Всем 10 заявкам назначены номера ТТН")

        expected_numbers = [
            str(int(START_NUMBER) - index) for index in range(ORDER_COUNT)
        ]
        actual_numbers = [order["ttn_number"] for order in assigned]
        report.check(
            actual_numbers == expected_numbers,
            f"Номера ТТН идут по порядку: {actual_numbers[0]}…{actual_numbers[-1]}")
        expected_rows = list(range(25, 15, -1))
        actual_rows = [order["order_row"] for order in assigned]
        report.check(
            actual_rows == expected_rows,
            f"Строки реестра: {actual_rows[0]}…{actual_rows[-1]}")

        registry_orders = orders_db.get_batch_orders(batch_id)
        registry_path = registry_filler.generate_registry_docx(
            registry_orders, "Самотест", "Самопроверка_реестр")
        registry_doc = Document(registry_path)
        table = registry_doc.tables[0]
        report.check(
            table.rows[25].cells[2].text == expected_numbers[0],
            "В строке 25 реестра стоит первая выбранная ТТН")
        report.check(
            table.rows[16].cells[2].text == expected_numbers[-1],
            "В строке 16 реестра стоит последняя выбранная ТТН")
        report.check(
            all(table.rows[row].cells[2].text.strip() == ""
                for row in range(1, 16)),
            "Строки 1–15 реестра остались пустыми")

        # Отдельная проверка атомарности: одна заявка уже назначена,
        # повторная попытка с ней должна откатить всю операцию.
        extra_one = orders_db.save_order(
            _order_data(11), master_id=MASTER_ID, driver_id=900011)
        extra_two = orders_db.save_order(
            _order_data(12), master_id=MASTER_ID, driver_id=900012)
        rollback_batch = orders_db.create_batch(MASTER_ID, "5555000")
        orders_db.assign_orders_to_batch(rollback_batch, [extra_one])
        rollback_worked = False
        try:
            orders_db.assign_orders_to_batch(
                rollback_batch, [extra_two, extra_one])
        except ValueError:
            rollback_worked = True
        report.check(rollback_worked, "Повторное назначение уже занятой ТТН отклонено")
        report.check(
            orders_db.get_order(extra_two)["status"] == "new"
            and orders_db.get_batch(rollback_batch)["counter"] == 1,
            "При ошибке назначение второй заявки откатилось")

        front_copy, back_copy, registry_copy = _copy_outputs(
            record, registry_path, output_dir)

        # PDF — только для визуальной самопроверки. В рабочем пакете
        # по-прежнему остаются только Word-файлы.
        soffice = find_soffice()
        try:
            import pymupdf
        except ImportError:
            pymupdf = None

        if not soffice or pymupdf is None:
            report.skip("Проверка PDF и пустых страниц: нет LibreOffice или PyMuPDF")
        else:
            try:
                front_pdf = _convert_to_pdf(soffice, front_copy, output_dir)
                back_pdf = _convert_to_pdf(soffice, back_copy, output_dir)
                registry_pdf = _convert_to_pdf(soffice, registry_copy, output_dir)
                front_doc_pdf = pymupdf.open(str(front_pdf))
                back_doc_pdf = pymupdf.open(str(back_pdf))
                registry_doc_pdf = pymupdf.open(str(registry_pdf))
                report.check(
                    front_doc_pdf.page_count == expected_sections,
                    f"Передний PDF: {front_doc_pdf.page_count} страниц "
                    f"(ожидалось {expected_sections})")
                report.check(
                    back_doc_pdf.page_count == expected_sections,
                    f"Задний PDF: {back_doc_pdf.page_count} страниц "
                    f"(ожидалось {expected_sections})")
                front_texts = [
                    front_doc_pdf[i].get_text("text")
                    for i in range(front_doc_pdf.page_count)
                ]
                back_texts = [
                    back_doc_pdf[i].get_text("text")
                    for i in range(back_doc_pdf.page_count)
                ]
                report.check(
                    all(text.strip() for text in front_texts),
                    "В переднем PDF нет пустых страниц")
                report.check(
                    all(text.strip() for text in back_texts),
                    "В заднем PDF нет пустых страниц")
                report.check(
                    all(
                        f"САМОТЕСТ-{number:02d}" in front_texts[(number - 1) * 4]
                        for number in range(1, ORDER_COUNT + 1)
                    ),
                    "Порядок заявок в переднем файле сохранён")
                report.check(
                    all(
                        f"Тестовый водитель {number}" in back_texts[(number - 1) * 4]
                        for number in range(1, ORDER_COUNT + 1)
                    ),
                    "Порядок заявок в заднем файле сохранён")
                report.check(
                    registry_doc_pdf.page_count >= 1
                    and registry_doc_pdf[0].get_text("text").strip(),
                    "Реестр открывается и содержит данные")
            except Exception as exc:
                report.check(False, "Проверка PDF", str(exc))

        report.info("")
        report.info("Файлы для просмотра:")
        report.info(f"  {front_copy}")
        report.info(f"  {back_copy}")
        report.info(f"  {registry_copy}")
        report.info("")
        report.info(
            "Проверено: 10 заявок, 4 экземпляра на заявку, "
            f"{ORDER_COUNT * 4} страниц в каждом файле, строки 25–16.")
        report.info("Реальные orders.db, data/ и реестр не изменялись.")
        report.info(
            "Word-файлы можно открыть и распечатать; PDF — только для "
            "визуальной проверки страниц.")
        report.info("")
        report.info("=" * 70)
        report.info(
            f"ИТОГ: {'УСПЕШНО' if report.failures == 0 else 'ЕСТЬ ОШИБКИ'}; "
            f"ошибок: {report.failures}, пропусков: {report.skips}")
        report.info("=" * 70)
        return report

    finally:
        if keep_temp:
            print(f"Временная папка сохранена: {runtime}")
        else:
            shutil.rmtree(runtime, ignore_errors=True)


def main():
    parser = argparse.ArgumentParser(
        description="Автономная самопроверка пакетной печати ТТН")
    parser.add_argument(
        "--template",
        help="Путь к шаблону ТТН (иначе берётся Шаблоны/ или демо-шаблон)")
    parser.add_argument(
        "--out",
        help="Куда положить образцы и отчёт "
             "(по умолчанию Результат/Самопроверка_пакетной_печати)")
    parser.add_argument(
        "--keep-temp", action="store_true",
        help="Оставить временную базу и промежуточные файлы")
    args = parser.parse_args()

    template = find_template(args.template)
    output_dir = Path(args.out) if args.out else (
        PROJECT_ROOT / "Результат" /
        f"Самопроверка_пакетной_печати_{datetime.now():%Y%m%d_%H%M%S}"
    )

    try:
        report = run_selftest(template, output_dir, args.keep_temp)
        report_path = output_dir / "Отчёт_самопроверки.txt"
        output_dir.mkdir(parents=True, exist_ok=True)
        report_path.write_text(report.text(), encoding="utf-8")
        print(report.text())
        print(f"\nОтчёт сохранён: {report_path}")
        return 0 if report.failures == 0 else 1
    finally:
        if not args.keep_temp:
            shutil.rmtree(SELFTEST_HOME, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
