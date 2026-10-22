# =============================================================
# ГЕНЕРАЦИЯ РЕЕСТРА ТТН (Word + PDF)
# =============================================================
import os
from docx import Document
from docx.shared import Pt, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RESULT_FOLDER = "Результат"
BATCH_SIZE = 25
COMPANY_SHORT_NAME = os.environ.get("COMPANY_SHORT_NAME", "Демо-Строй")

MONTHS_RU = {
    1: 'январь', 2: 'февраль', 3: 'март', 4: 'апрель', 5: 'май', 6: 'июнь',
    7: 'июль', 8: 'август', 9: 'сентябрь', 10: 'октябрь', 11: 'ноябрь', 12: 'декабрь',
}

# Ширины колонок, сумма = 19.0 см
COLUMNS = [
    ("№\nп/п", 1.0),
    ("Дата", 1.9),
    ("№\nнакладной", 2.3),
    ("Наименование\nобъекта", 5.2),
    ("Фамилия И.О.\nмастера", 2.6),
    ("Фамилия И.О.\nводителя", 3.4),
    ("№\nА-маш.", 2.6),
]

FONT_NAME = "Times New Roman"
FONT_SIZE_HEADER = Pt(12)
FONT_SIZE_DATA = Pt(12)

CM_TO_DXA = 566.93

ROW_HEIGHT_CM = 0.6
ROW_HEIGHT_RULE = 'atLeast'
HEADER_HEIGHT_CM = 1.2
HEADER_HEIGHT_RULE = 'atLeast'


def _set_table_fixed_layout(table, widths_cm):
    tbl = table._tbl
    tblPr = tbl.tblPr

    for el in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(el)
    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)

    for el in tblPr.findall(qn('w:tblW')):
        tblPr.remove(el)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(sum(widths_cm) * CM_TO_DXA)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)

    for grid in tbl.findall(qn('w:tblGrid')):
        tbl.remove(grid)
    grid = OxmlElement('w:tblGrid')
    for w in widths_cm:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(int(w * CM_TO_DXA)))
        grid.append(col)
    tbl.insert(list(tbl).index(tblPr) + 1, grid)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            for el in tcPr.findall(qn('w:tcW')):
                tcPr.remove(el)
            tcW = OxmlElement('w:tcW')
            tcW.set(qn('w:w'), str(int(widths_cm[i] * CM_TO_DXA)))
            tcW.set(qn('w:type'), 'dxa')
            tcPr.append(tcW)


def _set_row_height(row, height_cm, rule='atLeast'):
    trPr = row._tr.get_or_add_trPr()
    for el in trPr.findall(qn('w:trHeight')):
        trPr.remove(el)
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(height_cm * CM_TO_DXA)))
    trHeight.set(qn('w:hRule'), rule)
    trPr.append(trHeight)


def _set_cell_margins(cell, top=5, start=25, bottom=5, end=25):
    tcPr = cell._tc.get_or_add_tcPr()
    for el in tcPr.findall(qn('w:tcMar')):
        tcPr.remove(el)
    tcMar = OxmlElement('w:tcMar')
    for tag, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        el = OxmlElement(f'w:{tag}')
        el.set(qn('w:w'), str(value))
        el.set(qn('w:type'), 'dxa')
        tcMar.append(el)
    tcPr.append(tcMar)


def _apply_font(cell, size, bold=False):
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.name = FONT_NAME
            run.font.size = size
            run.bold = bold
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.append(rFonts)
            rFonts.set(qn('w:ascii'), FONT_NAME)
            rFonts.set(qn('w:hAnsi'), FONT_NAME)
            rFonts.set(qn('w:cs'), FONT_NAME)
            rFonts.set(qn('w:eastAsia'), FONT_NAME)


def _parse_date(date_str):
    try:
        parts = date_str.split('.')
        if len(parts) >= 3:
            return int(parts[0]), int(parts[1]), int(parts[2])
    except Exception:
        pass
    return None


def build_period_title(orders):
    months = [];
    year = None
    for o in orders:
        p = _parse_date(o.get('date', ''))
        if not p: continue
        _, month, yr = p
        if month not in months: months.append(month)
        if year is None: year = yr
    months.sort()
    if not months:
        period = ""
    elif len(months) == 1:
        period = f"за {MONTHS_RU[months[0]]}"
    elif len(months) == 2:
        period = f"за {MONTHS_RU[months[0]]} и {MONTHS_RU[months[1]]}"
    else:
        names = [MONTHS_RU[m] for m in months[:-1]]
        period = f"за {', '.join(names)} и {MONTHS_RU[months[-1]]}"
    if year is not None:
        period += f" 20{year:02d}г"
    return f"{COMPANY_SHORT_NAME} {period}".strip()


def generate_registry_docx(orders, master_name, output_name, manual_entries=None):
    if not os.path.exists(RESULT_FOLDER):
        os.makedirs(RESULT_FOLDER)

    doc = Document()

    for section in doc.sections:
        section.left_margin = Cm(0.7)
        section.right_margin = Cm(0.7)
        section.top_margin = Cm(0.8)
        section.bottom_margin = Cm(0.8)

    style = doc.styles['Normal']
    style.font.name = FONT_NAME
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), FONT_NAME)
    rfonts.set(qn('w:hAnsi'), FONT_NAME)
    rfonts.set(qn('w:cs'), FONT_NAME)

    # Заголовок «РЕЕСТР»
    p = doc.add_paragraph();
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0);
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("РЕЕСТР");
    r.bold = True;
    r.font.size = Pt(14);
    r.font.name = FONT_NAME

    # Подзаголовок
    p = doc.add_paragraph();
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Использованных товарно-транспортных накладных")
    r.font.size = Pt(12);
    r.font.name = FONT_NAME

    # Период
    all_orders_for_title = list(orders)
    if manual_entries:
        all_orders_for_title.extend(manual_entries)

    p = doc.add_paragraph();
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(build_period_title(all_orders_for_title))
    r.font.size = Pt(12);
    r.font.name = FONT_NAME

    # Создаем слоты (от 25 строки до 1)
    slots = [None] * BATCH_SIZE

    # 1. Сначала ставим ручные записи (у них есть target_row)
    if manual_entries:
        for me in manual_entries:
            row = me.get('target_row')
            if row and 1 <= row <= BATCH_SIZE:
                idx = BATCH_SIZE - row
                slots[idx] = me

    # 2. Затем ставим реальные заявки (они вытесняются, если место занято)
    sorted_orders = sorted(orders, key=lambda x: x.get('order_row', 0))
    for o in sorted_orders:
        desired_row = o.get('order_row', 1)
        idx = BATCH_SIZE - desired_row
        # Ищем свободное место вверх (в сторону увеличения строки)
        while idx >= 0 and slots[idx] is not None:
            idx -= 1

        if idx >= 0:
            slots[idx] = o

    table = doc.add_table(rows=1 + BATCH_SIZE, cols=len(COLUMNS))
    table.style = 'Table Grid'
    table.autofit = False
    table.allow_autofit = False

    # Шапка
    hdr = table.rows[0].cells
    for i, (label, _) in enumerate(COLUMNS):
        hdr[i].text = ""
        cell_p = hdr[i].paragraphs[0]
        cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_p.paragraph_format.space_before = Pt(0)
        cell_p.paragraph_format.space_after = Pt(0)
        cell_p.paragraph_format.line_spacing = 1.0
        parts = label.split('\n')
        for j, part in enumerate(parts):
            if j > 0:
                cell_p.add_run().add_break()
            cell_p.add_run(part)
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _apply_font(hdr[i], FONT_SIZE_HEADER, bold=True)
        _set_cell_margins(hdr[i])
    _set_row_height(table.rows[0], HEADER_HEIGHT_CM, HEADER_HEIGHT_RULE)

    # Данные
    for row_idx in range(1, BATCH_SIZE + 1):
        cells = table.rows[row_idx].cells
        idx = BATCH_SIZE - row_idx
        entry = slots[idx]

        for c in cells:
            c.text = ""
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        if entry:
            values = [
                str(row_idx),
                entry.get('date', ''),
                entry.get('ttn_number', ''),
                entry.get('loading_point', ''),
                entry.get('master_name', master_name),
                entry.get('driver_name', ''),
                entry.get('car_plate', ''),
            ]
        else:
            values = [str(row_idx), "", "", "", "", "", ""]

        for i, val in enumerate(values):
            cell_p = cells[i].paragraphs[0]
            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell_p.paragraph_format.space_before = Pt(0)
            cell_p.paragraph_format.space_after = Pt(0)
            cell_p.paragraph_format.line_spacing = 1.0
            cell_p.add_run(str(val))

        for c in cells:
            _apply_font(c, FONT_SIZE_DATA)
            _set_cell_margins(c)

        _set_row_height(table.rows[row_idx], ROW_HEIGHT_CM, ROW_HEIGHT_RULE)

    widths = [c[1] for c in COLUMNS]
    _set_table_fixed_layout(table, widths)

    # Подпись
    sign = doc.add_paragraph()
    sign.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sign.paragraph_format.space_before = Pt(10)
    sign.paragraph_format.space_after = Pt(0)
    r = sign.add_run(f"Мастер {master_name}")
    r.font.size = Pt(12);
    r.font.name = FONT_NAME

    path = os.path.join(RESULT_FOLDER, f"{output_name}.docx")
    doc.save(path)
    print(f"✅ Реестр DOCX создан: {path}")
    return path


def generate_registry(orders, master_name, output_name, manual_entries=None):
    from word_filler import convert_to_pdf
    docx_path = generate_registry_docx(orders, master_name, output_name, manual_entries)
    ok, result = convert_to_pdf(docx_path)
    pdf_path = result if ok else None
    return docx_path, pdf_path