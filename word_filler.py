# =============================================================
# ЗАПОЛНЕНИЕ WORD И КОНВЕРТАЦИЯ В PDF (Финальная версия с выравниванием)
# =============================================================
import os
import subprocess
import shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy
from docx.table import _Row
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    import pymupdf as fitz
    PYMUPDF_OK = True
except ImportError:
    PYMUPDF_OK = False

TEMPLATE = "Шаблоны/Шаблон_ТТН.docx"
DEMO_TEMPLATE = "assets/Шаблон_ТТН_демо.docx"
RESULT_FOLDER = "Результат"
COMPANY_NAME = os.environ.get("COMPANY_NAME", "ООО «Демо-Строй»")
RELEASE_ALLOWED_BY = os.environ.get("RELEASE_ALLOWED_BY", "Руководитель")

MONTHS_RU = {
    '01': 'января', '02': 'февраля', '03': 'марта',
    '04': 'апреля', '05': 'мая', '06': 'июня',
    '07': 'июля', '08': 'августа', '09': 'сентября',
    '10': 'октября', '11': 'ноября', '12': 'декабря',
}

def month_to_text(month_str):
    if not month_str: return ''
    m = str(month_str).strip().zfill(2)
    return MONTHS_RU.get(m, str(month_str))

def clean_driver_name(name):
    name = (name or '').strip()
    if name.lower().startswith('водитель'):
        name = name[len('водитель'):].strip()
    return name


def _replace_placeholders(paragraph, data):
    """Заменяет {{ключи}} в абзаце, не разрушая табуляцию и форматирование.

    Placeholder может быть разбит на несколько run'ов. Поэтому работаем
    с текстовыми XML-узлами w:t, а не с paragraph.text: так сохраняются
    элементы w:tab (они нужны для позиционирования полей на бланке).
    """
    nodes = paragraph._p.xpath('.//w:t')
    if not nodes:
        return

    for key, value in data.items():
        if key == 'cargos':
            continue
        placeholder = "{{" + key + "}}"
        replacement = str(value)

        while True:
            texts = [node.text or '' for node in nodes]
            full = ''.join(texts)
            start = full.find(placeholder)
            if start < 0:
                break
            end = start + len(placeholder)

            affected = []
            pos = 0
            for node in nodes:
                text = node.text or ''
                nstart, nend = pos, pos + len(text)
                if nend > start and nstart < end:
                    local_start = max(start, nstart) - nstart
                    local_end = min(end, nend) - nstart
                    affected.append((node, text, local_start, local_end))
                pos = nend

            if not affected:
                break

            first_node, first_text, first_start, first_end = affected[0]
            if len(affected) == 1:
                first_node.text = (
                    first_text[:first_start] + replacement + first_text[first_end:]
                )
            else:
                first_node.text = first_text[:first_start] + replacement
                for node, text, _, local_end in affected[1:]:
                    node.text = text[local_end:]


def fill_documents(data, output_name):
    if not os.path.exists(RESULT_FOLDER):
        os.makedirs(RESULT_FOLDER)

    defaults = {
        'day': '', 'month': '', 'year': '', 'waybill_number': '',
        'car_model': '', 'car_plate': '', 'trailer_info': '-',
        'driver_name': '', 'driver_name_role': '',
        'customer': COMPANY_NAME,
        'shipper': COMPANY_NAME,
        'consignee': COMPANY_NAME,
        'release_basis': 'Распоряжение директора',
        'loading_point': '', 'unloading_point': '',
        'release_allowed_by': RELEASE_ALLOWED_BY,
        'shipper_handed_by': '', 'consignee_received_by': '',
        'total_mass_text': '', 'total_places_text': '',
        'total_vat_text': '', 'total_vat_cents': '',
        'total_with_vat_text': '', 'total_with_vat_cents': '',
        'shipper_unp': '', 'consignee_unp': '', 'customer_unp': ''
    }
    for key, value in defaults.items():
        data.setdefault(key, value)

    # В старых заявках и шаблонах копейки могли оставаться "00" даже без
    # рассчитанной суммы. Если суммы прописью нет, поле копеек делаем пустым.
    for text_key, cents_key in (
        ('total_vat_text', 'total_vat_cents'),
        ('total_with_vat_text', 'total_with_vat_cents'),
    ):
        if not str(data.get(text_key) or '').strip():
            data[cents_key] = ''

    data['month'] = month_to_text(data.get('month', ''))
    driver_clean = clean_driver_name(data.get('driver_name', ''))
    data['driver_name'] = driver_clean
    data['driver_name_role'] = f"Водитель {driver_clean}" if driver_clean else ''

    template_path = TEMPLATE if os.path.exists(TEMPLATE) else DEMO_TEMPLATE
    if not os.path.exists(template_path):
        raise FileNotFoundError(
            "Не найден шаблон ТТН. Запустите python create_template.py "
            "или положите файл в Шаблоны/Шаблон_ТТН.docx"
        )
    doc = Document(template_path)

    cargos = data.get('cargos', [])
    cargo_table_processed = False

    for paragraph in doc.paragraphs:
        _replace_placeholders(paragraph, data)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    is_cents = any(
                        "{{" + key + "}}" in p.text
                        for key in ('total_vat_cents', 'total_with_vat_cents')
                    )
                    _replace_placeholders(p, data)
                    if is_cents:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.text = p.text.strip()

        if not cargo_table_processed:
            cargo_row = None
            row_idx = -1
            for idx, row in enumerate(table.rows):
                if any("cargo.cargo_name" in cell.text for cell in row.cells):
                    cargo_row = row
                    row_idx = idx
                    break

            if cargo_row:
                template_tr = cargo_row._tr
                tbl = table._tbl
                tbl.remove(template_tr)

                for i, cargo in enumerate(cargos):
                    new_tr = deepcopy(template_tr)
                    trPr = new_tr.find(qn('w:trPr'))
                    if trPr is None:
                        trPr = OxmlElement('w:trPr')
                        new_tr.insert(0, trPr)
                    for el in trPr.findall(qn('w:trHeight')):
                        trPr.remove(el)
                    trHeight = OxmlElement('w:trHeight')
                    trHeight.set(qn('w:val'), '200')
                    trHeight.set(qn('w:hRule'), 'atLeast')
                    trPr.append(trHeight)

                    # Шаблон ожидает si (ед. изм.) и unit (количество).
                    # В данных бота unit — единица измерения, quantity — количество.
                    cargo_values = dict(cargo)
                    if 'si' not in cargo_values:
                        cargo_values['si'] = cargo.get('unit', '')
                    if 'quantity' in cargo_values:
                        cargo_values['unit'] = cargo.get('quantity', '')

                    new_row = _Row(new_tr, table)
                    for cell in new_row.cells:
                        tcPr = cell._tc.get_or_add_tcPr()
                        tcMar = tcPr.find(qn('w:tcMar'))
                        if tcMar is None:
                            tcMar = OxmlElement('w:tcMar')
                            tcPr.append(tcMar)
                        for side in ['top', 'bottom', 'start', 'end', 'left', 'right']:
                            node = tcMar.find(qn(f'w:{side}'))
                            if node is None:
                                node = OxmlElement(f'w:{side}')
                                tcMar.append(node)
                            node.set(qn('w:w'), '0')
                            node.set(qn('w:type'), 'dxa')

                        for p in cell.paragraphs:
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(0)
                            p.paragraph_format.line_spacing = 1.0
                            _replace_placeholders(
                                p, {f"cargo.{k}": v for k, v in cargo_values.items()})
                            if len(cargos) > 5:
                                font_sz = Pt(7) if len(cargos) > 8 else Pt(8)
                                for r in p.runs:
                                    r.font.size = font_sz
                    tbl.insert(row_idx + i, new_tr)
                cargo_table_processed = True

    result_path = os.path.join(RESULT_FOLDER, f"{output_name}_ТТН.docx")
    doc.save(result_path)
    return result_path, result_path, result_path

def find_soffice():
    candidates = [
        shutil.which("soffice"), shutil.which("soffice.exe"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice", "/usr/local/bin/soffice"
    ]
    for p in candidates:
        if p and os.path.exists(p): return p
    return None

def convert_to_pdf(docx_path):
    if not os.path.exists(docx_path): return False, "Файл не найден"
    docx_path = os.path.abspath(docx_path)
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    out_dir = os.path.dirname(docx_path)
    if os.path.exists(pdf_path):
        try: os.remove(pdf_path)
        except: pass
    soffice = find_soffice()
    if not soffice: return False, "LibreOffice не найден"
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
                       check=True, capture_output=True, timeout=180)
        if not os.path.exists(pdf_path):
            return False, "PDF не создан"
        return True, pdf_path
    except Exception as e:
        return False, str(e)
