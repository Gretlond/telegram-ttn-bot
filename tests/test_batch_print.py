import os
import shutil
import subprocess
from types import SimpleNamespace

import pytest
from docx import Document

import batch_word
import orders_db
import word_filler


@pytest.fixture
def temp_db(tmp_path, monkeypatch):
    monkeypatch.setattr(orders_db, "DB_FILE", str(tmp_path / "test_batch.db"))
    orders_db.init_db()
    return orders_db


def _order_data(number):
    return {
        "day": "22", "month": "10", "year": "26",
        "waybill_number": f"WB-{number}",
        "car_model": "МАЗ", "car_plate": f"1234 AB-{number}",
        "trailer_info": "-",
        "loading_point": "Склад", "unloading_point": "Объект",
        "driver_name": f"Водитель {number}",
        "cargos": [{
            "cargo_name": f"Груз {number}",
            "si": "т", "unit": "10", "cargo_qty": "10",
        }],
    }


def test_assign_orders_continues_batch_counter(temp_db):
    first = temp_db.save_order(_order_data(1), master_id=500, driver_id=900)
    second = temp_db.save_order(_order_data(2), master_id=500, driver_id=901)
    batch_id = temp_db.create_batch(500, "0217975")

    assignments = temp_db.assign_orders_to_batch(batch_id, [first, second])

    assert [a["ttn_number"] for a in assignments] == ["0217975", "0217974"]
    assert [a["order_row"] for a in assignments] == [25, 24]
    assert temp_db.get_order(first)["status"] == "printed"
    assert temp_db.get_order(second)["status"] == "printed"
    assert temp_db.get_batch(batch_id)["counter"] == 2


def test_assign_orders_rolls_back_when_one_order_changed(temp_db):
    first = temp_db.save_order(_order_data(1), master_id=500, driver_id=900)
    second = temp_db.save_order(_order_data(2), master_id=500, driver_id=901)
    batch_id = temp_db.create_batch(500, "0217975")
    temp_db.assign_orders_to_batch(batch_id, [first])

    with pytest.raises(ValueError):
        temp_db.assign_orders_to_batch(batch_id, [second, first])

    assert temp_db.get_order(second)["status"] == "new"
    assert temp_db.get_order(second)["ttn_number"] is None
    assert temp_db.get_batch(batch_id)["counter"] == 1


def _make_order_doc(output_dir, number):
    path, _, _ = word_filler.fill_documents(
        _order_data(number), f"batch_order_{number}", output_dir=str(output_dir))
    return path


def _section_layout(doc):
    return [
        (s.page_width, s.page_height, s.left_margin, s.right_margin,
         s.top_margin, s.bottom_margin, s.orientation)
        for s in doc.sections
    ]


def test_batch_word_repeats_each_order_four_times(tmp_path, monkeypatch):
    monkeypatch.setattr(word_filler, "RESULT_FOLDER", str(tmp_path / "result"))
    docs = [
        _make_order_doc(tmp_path, 1),
        _make_order_doc(tmp_path, 2),
    ]
    source_layout = _section_layout(Document(docs[0]))

    front, back = batch_word.build_batch_documents(
        docs, tmp_path / "batch", "test_batch")

    front_doc = Document(front)
    back_doc = Document(back)
    assert len(front_doc.sections) == 2 * batch_word.COPIES_PER_TTN
    assert len(back_doc.sections) == 2 * batch_word.COPIES_PER_TTN
    assert _section_layout(front_doc) == source_layout[:1] * 8
    assert _section_layout(back_doc) == source_layout[1:2] * 8


def test_batch_word_pdf_has_no_blank_pages(tmp_path, monkeypatch):
    try:
        import pymupdf
    except ImportError:
        pytest.skip("PyMuPDF не установлен")

    soffice = shutil.which("soffice") or shutil.which("soffice.exe")
    if not soffice:
        for candidate in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            "/usr/bin/soffice",
        ):
            if os.path.exists(candidate):
                soffice = candidate
                break
    if not soffice:
        pytest.skip("LibreOffice не установлен")

    monkeypatch.setattr(word_filler, "RESULT_FOLDER", str(tmp_path / "result"))
    docs = [
        _make_order_doc(tmp_path, 1),
        _make_order_doc(tmp_path, 2),
    ]
    front, back = batch_word.build_batch_documents(
        docs, tmp_path / "batch", "test_batch_pdf")

    for docx_path in (front, back):
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf",
             "--outdir", str(tmp_path / "batch"), docx_path],
            check=True, capture_output=True, timeout=180,
        )
    for path, expected_first in (
        (front.replace(".docx", ".pdf"), "WB-1"),
        (back.replace(".docx", ".pdf"), "Водитель 1"),
    ):
        pdf = pymupdf.open(path)
        assert pdf.page_count == 8
        texts = [pdf[i].get_text("text") for i in range(pdf.page_count)]
        assert all(text.strip() for text in texts)
        assert all(expected_first in text for text in texts[:4])
        assert all(expected_first not in text for text in texts[4:])


def test_generate_batch_creates_two_side_files(
    temp_db, tmp_path, monkeypatch
):
    import batch_print

    first = temp_db.save_order(_order_data(1), master_id=500, driver_id=900)
    second = temp_db.save_order(_order_data(2), master_id=500, driver_id=901)
    assignments = batch_print._make_assignments([first, second], "0217975", 0)
    plan = {
        "selected": [first, second],
        "start_number": "0217975",
        "batch_id": None,
        "counter": 0,
        "assignments": assignments,
    }
    monkeypatch.setattr(batch_print, "PENDING_FILE", str(tmp_path / "pending.json"))
    monkeypatch.setattr(word_filler, "RESULT_FOLDER", str(tmp_path / "result"))

    record = batch_print._generate_batch(500, plan)

    assert Document(record["front_path"]).sections
    assert Document(record["back_path"]).sections
    assert len(Document(record["front_path"]).sections) == 8
    assert len(Document(record["back_path"]).sections) == 8
    assert batch_print._get_pending(500)["front_path"] == record["front_path"]


def test_batch_registry_keeps_selection_order_from_row_25(
    temp_db, tmp_path, monkeypatch
):
    import registry_filler

    first = temp_db.save_order(_order_data(1), master_id=500, driver_id=900)
    second = temp_db.save_order(_order_data(2), master_id=500, driver_id=901)
    batch_id = temp_db.create_batch(500, "0217975")
    temp_db.assign_orders_to_batch(batch_id, [first, second])

    monkeypatch.setattr(registry_filler, "RESULT_FOLDER", str(tmp_path))
    path = registry_filler.generate_registry_docx(
        temp_db.get_batch_orders(batch_id), "Мастер", "batch_registry")
    table = Document(path).tables[0]

    assert table.rows[25].cells[2].text == "0217975"
    assert table.rows[24].cells[2].text == "0217974"
    assert table.rows[25].cells[0].text == "25"
    assert table.rows[24].cells[0].text == "24"


def _call(user_id, data):
    return SimpleNamespace(
        id="callback-batch",
        data=data,
        message=SimpleNamespace(chat=SimpleNamespace(id=user_id), message_id=10),
    )


def _pending_record(tmp_path, batch_id, assignments, start_number="0217975", counter=0):
    front = tmp_path / "front.docx"
    back = tmp_path / "back.docx"
    front.write_bytes(b"front")
    back.write_bytes(b"back")
    return {
        "created_at": "2026-10-28 12:00:00",
        "assignments": assignments,
        "batch_id": batch_id,
        "counter": counter,
        "start_number": start_number,
        "front_path": str(front),
        "back_path": str(back),
    }


def test_batch_menu_renders_available_orders(temp_db, tmp_path, monkeypatch):
    import batch_print

    temp_db.save_order(_order_data(1), master_id=500, driver_id=900)
    monkeypatch.setattr(batch_print, "PENDING_FILE", str(tmp_path / "pending.json"))
    batch_print._selection_state.clear()
    messages = []
    monkeypatch.setattr(batch_print.bot, "send_message",
                        lambda *a, **k: messages.append((a, k)))

    batch_print.render_batch_menu(500)

    assert messages
    assert "Пакетная печать" in messages[0][0][1]
    markup = messages[0][1]["reply_markup"]
    assert any(
        button.callback_data == "batch_toggle_1"
        for row in markup.keyboard for button in row
    )


def test_batch_confirmation_updates_registry_and_clears_pending(
    temp_db, tmp_path, monkeypatch
):
    import batch_print

    first = temp_db.save_order(_order_data(1), master_id=500, driver_id=900)
    second = temp_db.save_order(_order_data(2), master_id=500, driver_id=901)
    batch_id = temp_db.create_batch(500, "0217975")
    assignments = batch_print._make_assignments([first, second], "0217975", 0)
    record = _pending_record(tmp_path, batch_id, assignments)

    monkeypatch.setattr(batch_print, "PENDING_FILE", str(tmp_path / "pending.json"))
    batch_print._put_pending(500, record)
    messages = []
    monkeypatch.setattr(batch_print.bot, "answer_callback_query", lambda *a, **k: None)
    monkeypatch.setattr(batch_print.bot, "send_message",
                        lambda *a, **k: messages.append((a, k)))

    batch_print.batch_confirm_print(_call(500, "batch_confirm_print"))

    assert temp_db.get_order(first)["status"] == "printed"
    assert temp_db.get_order(first)["order_row"] == 25
    assert temp_db.get_order(second)["order_row"] == 24
    assert temp_db.get_batch(batch_id)["counter"] == 2
    assert batch_print._get_pending(500) is None
    assert any("Печать подтверждена" in args[1] for args, _ in messages)


def test_batch_cancel_deletes_files_and_keeps_orders_new(
    temp_db, tmp_path, monkeypatch
):
    import batch_print

    first = temp_db.save_order(_order_data(1), master_id=500, driver_id=900)
    batch_id = temp_db.create_batch(500, "0217975")
    assignments = batch_print._make_assignments([first], "0217975", 0)
    record = _pending_record(tmp_path, batch_id, assignments)

    monkeypatch.setattr(batch_print, "PENDING_FILE", str(tmp_path / "pending.json"))
    batch_print._put_pending(500, record)
    monkeypatch.setattr(batch_print.bot, "answer_callback_query", lambda *a, **k: None)
    monkeypatch.setattr(batch_print.bot, "send_message", lambda *a, **k: None)

    batch_print.batch_cancel_pending(_call(500, "batch_cancel_pending"))

    assert temp_db.get_order(first)["status"] == "new"
    assert temp_db.get_batch(batch_id)["counter"] == 0
    assert batch_print._get_pending(500) is None
    assert not (tmp_path / "front.docx").exists()
    assert not (tmp_path / "back.docx").exists()
