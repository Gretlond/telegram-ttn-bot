from docx import Document


def test_money_to_words_parts():
    from utils import amount_to_words_parts

    assert amount_to_words_parts(350.40) == ("триста пятьдесят", "40")
    assert amount_to_words_parts(61.5) == ("шестьдесят один", "50")
    assert amount_to_words_parts(369) == ("триста шестьдесят девять", "00")
    assert amount_to_words_parts(0) == ("", "")
    assert amount_to_words_parts(0.5) == ("ноль", "50")


def test_calculate_vat_totals_uses_row_totals():
    from utils import calculate_vat_totals

    cargos = [
        {"total_cost": "307.50", "vat_rate": "20%", "vat_cost": "61.50",
         "total_with_vat": "369.00"},
        {"total_cost": 100, "vat_rate": "Без НДС", "vat_cost": 0,
         "total_with_vat": 100},
    ]
    assert calculate_vat_totals(cargos) == (407.5, 61.5, 469.0)


def test_calculate_vat_totals_falls_back_to_qty_and_price():
    from utils import calculate_vat_totals

    cargos = [{"cargo_qty": "10", "price": "100", "vat_rate": "20%"}]
    assert calculate_vat_totals(cargos) == (1000.0, 200.0, 1200.0)


def _cargo_row(doc, name):
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if name in values:
                return values
    raise AssertionError(f"Строка груза {name!r} не найдена")


def _total_row(doc, first_words):
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if values and values[0] == first_words:
                return values
    raise AssertionError(f"Строка итогов {first_words!r} не найдена")


def test_word_cargo_columns_and_vat_totals():
    import word_filler
    from utils import amount_to_words_parts

    cargo = {
        "cargo_name": "Pesok",
        "unit": "тн",
        "quantity": "11",
        "cargo_qty": "20.5",
        "places": "11",
        "price": "15",
        "total_cost": "307.50",
        "vat_rate": "20%",
        "vat_cost": "61.50",
        "total_with_vat": "369.00",
    }
    vat_text, vat_cents = amount_to_words_parts(61.5)
    total_text, total_cents = amount_to_words_parts(369)
    data = {
        "waybill_number": "TEST",
        "day": "22", "month": "10", "year": "26",
        "car_model": "MAZ", "car_plate": "AA", "trailer_info": "-",
        "driver_name": "Ivanov", "loading_point": "A", "unloading_point": "B",
        "total_vat_text": vat_text,
        "total_vat_cents": vat_cents,
        "total_with_vat_text": total_text,
        "total_with_vat_cents": total_cents,
        "cargos": [cargo],
    }

    path, _, _ = word_filler.fill_documents(data, "pytest_ttn")
    doc = Document(path)

    row = _cargo_row(doc, "Pesok")
    assert row[0] == "Pesok"
    assert row[1] == "тн"
    assert row[2] == "11"
    assert all("{{cargo." not in value for value in row)

    vat_row = _total_row(doc, "шестьдесят один")
    assert vat_row[:3] == ["шестьдесят один"] * 3
    assert vat_row[3] == "50"

    total_row = _total_row(doc, "триста шестьдесят девять")
    assert total_row[:3] == ["триста шестьдесят девять"] * 3
    assert total_row[3] == "00"

    for table in doc.tables:
        for r in table.rows:
            for cell in r.cells:
                assert "{{cargo." not in cell.text


def test_word_keeps_cents_blank_without_calculation():
    import word_filler
    from utils import amount_to_words_parts

    assert amount_to_words_parts(0) == ("", "")

    cargo = {
        "cargo_name": "Pesok",
        "si": "тн",
        "unit": "20",
        "cargo_qty": "20",
    }
    path, _, _ = word_filler.fill_documents(
        {"waybill_number": "TEST-NO-COST", "cargos": [cargo]},
        "pytest_ttn_no_cost")
    doc = Document(path)

    cells = [
        cell.text.strip()
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    ]
    assert "00" not in cells


def test_word_header_fills_date_and_keeps_unloading_separate():
    import word_filler
    from docx import Document

    path, _, _ = word_filler.fill_documents(
        {"day": "22", "month": "10", "year": "26",
         "car_model": "MAZ", "car_plate": "AA", "waybill_number": "123456",
         "release_basis": "RELEASE", "loading_point": "LOAD",
         "unloading_point": "UNLOAD"},
        "pytest_ttn_header")

    doc = Document(path)
    rows = [
        [cell.text.strip() for cell in row.cells]
        for table in doc.tables
        for row in table.rows
    ]
    release_index = next(i for i, values in enumerate(rows) if "RELEASE" in values)
    loading_index = next(i for i, values in enumerate(rows) if "LOAD" in values)
    unloading_index = next(i for i, values in enumerate(rows) if "UNLOAD" in values)

    assert release_index == loading_index
    assert "UNLOAD" not in rows[release_index]
    assert unloading_index == release_index + 1

    text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "22" in text and "октября" in text and "26" in text
    assert "{{" not in text


def test_replace_placeholders_keeps_tab_elements():
    import word_filler
    from docx import Document

    doc = Document()
    paragraph = doc.add_paragraph()
    first = paragraph.add_run()
    first.add_tab()
    paragraph.add_run("{{value}}")

    word_filler._replace_placeholders(paragraph, {"value": "42"})

    assert paragraph._p.xpath(".//w:tab")
    assert paragraph.text.replace("\t", "") == "42"


def test_word_supports_legacy_si_and_unit_values():
    import word_filler

    cargo = {
        "cargo_name": "Pesok",
        "si": "тн",
        "unit": "11",
        "cargo_qty": "20.5",
        "places": "11",
    }
    path, _, _ = word_filler.fill_documents(
        {"waybill_number": "TEST2", "cargos": [cargo]}, "pytest_ttn_legacy")
    doc = Document(path)
    row = _cargo_row(doc, "Pesok")
    assert row[1] == "тн"
    assert row[2] == "11"


def test_trailer_info_prints_only_plate_number():
    import word_filler

    assert word_filler.clean_trailer_info("ГКБ 8352 (АК 5678-9)") == "АК 5678-9"
    assert word_filler.clean_trailer_info("А 8584 B-4 (А 8584 B-4)") == "А 8584 B-4"
    assert word_filler.clean_trailer_info("АК 5678-9") == "АК 5678-9"
    assert word_filler.clean_trailer_info("-") == "-"

    path, _, _ = word_filler.fill_documents(
        {"waybill_number": "TEST-TRAILER",
         "trailer_info": "ГКБ 8352 (АК 5678-9)"},
        "pytest_ttn_trailer")
    doc = Document(path)
    text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "АК 5678-9" in text
    assert "ГКБ 8352" not in text


def test_ttn_header_settings_are_used_in_document(monkeypatch):
    import database
    import word_filler

    monkeypatch.setattr(database, "TTN_HEADER", {
        "shipper": "ОАО «Тест-Отправитель»",
        "shipper_unp": "111111111",
        "consignee": "ОАО «Тест-Получатель»",
        "consignee_unp": "222222222",
        "customer": "ОАО «Тест-Заказчик»",
        "customer_unp": "333333333",
    })

    path, _, _ = word_filler.fill_documents(
        {"waybill_number": "TEST-HEADER"}, "pytest_ttn_header")
    doc = Document(path)

    top = doc.tables[0]
    assert len(top.rows) == 2
    assert [cell.text for cell in top.rows[0].cells] == [
        "", "Грузоотправитель", "Грузополучатель",
        "Заказчик автомобильной перевозки (плательщик)",
    ]
    assert [cell.text for cell in top.rows[1].cells] == [
        "УНП", "111111111", "222222222", "333333333",
    ]

    body = doc.tables[1]
    assert body.rows[4].cells[0].text == "ОАО «Тест-Заказчик»"
    assert body.rows[5].cells[0].text == "ОАО «Тест-Отправитель»"
    assert body.rows[6].cells[0].text == "ОАО «Тест-Получатель»"
