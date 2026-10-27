from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def _totals_table(doc):
    for table in doc.tables:
        if any(
            "total_vat_text" in cell.text
            for row in table.rows
            for cell in row.cells
        ):
            return table
    raise AssertionError("Таблица итогов не найдена")


def test_demo_template_keeps_totals_on_second_page():
    """Итоги и подписи не должны всплывать над датой на первой странице."""
    template = Path(__file__).resolve().parents[1] / "assets" / "Шаблон_ТТН_демо.docx"
    doc = Document(template)
    table = _totals_table(doc)

    # Раньше таблица была floating-таблицей с привязкой к верху страницы.
    assert table._tbl.tblPr.find(qn("w:tblpPr")) is None

    previous = table._tbl.getprevious()
    assert previous is not None
    assert previous.tag == qn("w:p")
    p_pr = previous.find(qn("w:pPr"))
    assert p_pr is not None
    # Итоги начинаются на второй странице через разрыв раздела, чтобы у
    # второй страницы был свой (меньший) верхний отступ.
    assert p_pr.find(qn("w:sectPr")) is not None
    assert len(doc.sections) == 2
    assert doc.sections[1].top_margin < doc.sections[0].top_margin


def test_demo_totals_rows_match_paper_form():
    template = Path(__file__).resolve().parents[1] / "assets" / "Шаблон_ТТН_демо.docx"
    doc = Document(template)
    table = _totals_table(doc)

    assert len(table.rows) == 5
    shipper_row = next(
        row
        for row in table.rows
        if any("shipper_handed_by" in cell.text for cell in row.cells)
    )
    assert any("consignee_received_by" in cell.text for cell in shipper_row.cells)


def test_demo_cargo_table_fits_content_width_and_centers_values():
    template = Path(__file__).resolve().parents[1] / "assets" / "Шаблон_ТТН_демо.docx"
    doc = Document(template)
    section = doc.sections[0]
    content_width = section.page_width - section.left_margin - section.right_margin

    table = next(
        table
        for table in doc.tables
        if any("{{cargo.cargo_name}}" in cell.text for row in table.rows for cell in row.cells)
    )
    cargo_row = next(
        row for row in table.rows
        if any("{{cargo.cargo_name}}" in cell.text for cell in row.cells)
    )

    assert abs(sum(column.width for column in table.columns) - content_width) < 10000
    assert table._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")) == "0"

    alignments = [cell.paragraphs[0].alignment for cell in cargo_row.cells]
    assert alignments[0] == WD_ALIGN_PARAGRAPH.LEFT
    assert all(alignment == WD_ALIGN_PARAGRAPH.CENTER for alignment in alignments[1:])
